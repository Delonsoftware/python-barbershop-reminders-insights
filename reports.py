"""
Output writers: CSVs for importing into an SMS/email tool, and two Excel workbooks
(daily messages, and the weekly business report with charts and a busiest-hours heatmap).
"""

import csv
import math
from datetime import timedelta
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from insights import money, pct_change

HEADER_FILL = PatternFill("solid", fgColor="1F2A44")
HEADER_FONT = Font(bold=True, color="FFFFFF")
GOOD_FONT = Font(color="1E7B34", bold=True)
BAD_FONT = Font(color="C00000", bold=True)
RISK_FILL = PatternFill("solid", fgColor="FCE4D6")
TIER_FILLS = {"Due": PatternFill("solid", fgColor="E2F0D9"),
              "Overdue": PatternFill("solid", fgColor="FFF2CC"),
              "Lapsed": PatternFill("solid", fgColor="F8CBAD")}
MONEY_FMT = '"$"#,##0.00'
CHANGE_FMT = "+0.0%;-0.0%;0.0%"


def write_csv(path: Path, rows: list, columns: list) -> None:
    """columns: list of (header, key)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow([header for header, _ in columns])
        for row in rows:
            writer.writerow([row[key] for _, key in columns])


def write_sheet(ws, headers, rows, widths, formats=None, wrap_cols=(), empty_note=None):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append([None if (isinstance(v, float) and math.isnan(v)) else v for v in row])
    for idx, name in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]
        for cells in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            if formats and name in formats:
                cells[0].number_format = formats[name]
            cells[0].alignment = Alignment(wrap_text=name in wrap_cols, vertical="top")
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    elif empty_note:
        ws["A2"] = empty_note


def fit_landscape(ws) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


# --------------------------------------------------------------------------
# Daily messages workbook
# --------------------------------------------------------------------------
REMINDER_COLUMNS = [("Time", "time_text"), ("Client", "client_name"), ("Barber", "barber"),
                    ("Service", "service"), ("Channel", "channel"), ("To", "to"), ("Risk", "risk"),
                    ("Message", "message"), ("Characters", "chars"), ("SMS Segments", "segments")]
REBOOK_COLUMNS = [("Tier", "tier"), ("Client ID", "client_id"), ("Client", "client_name"),
                  ("Channel", "channel"), ("To", "to"), ("Last Visit", "last_visit"),
                  ("Days Since", "days_since"), ("Usual Gap (days)", "usual_interval_days"),
                  ("Visits", "visits"), ("Usual Barber", "usual_barber"),
                  ("Usual Service", "usual_service"), ("Avg Spend", "avg_spend"),
                  ("Message", "message"), ("Characters", "chars"), ("SMS Segments", "segments")]


def write_messages_csvs(output_dir: Path, target, today, reminders: list, rebooking: list) -> tuple:
    for r in reminders:
        r["time_text"] = r["start"].strftime("%Y-%m-%d %H:%M")
    reminders_csv = output_dir / f"reminders_{target:%Y-%m-%d}.csv"
    rebooking_csv = output_dir / f"rebooking_{today:%Y-%m-%d}.csv"
    write_csv(reminders_csv, reminders, REMINDER_COLUMNS)
    write_csv(rebooking_csv, rebooking, REBOOK_COLUMNS)
    return reminders_csv, rebooking_csv


def write_messages_workbook(path: Path, target, today, reminders, skipped_reminders,
                            rebooking, skip_counts, config) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    reminders_ws, rebook_ws, skipped_ws = wb.create_sheet("Reminders"), wb.create_sheet("Rebooking"), wb.create_sheet("Not Messaged")

    summary["A1"] = f"{config['shop']['name']}: messages drafted {today:%b %d, %Y}"
    summary["A1"].font = Font(bold=True, size=15)
    summary["A2"] = "Drafts only: nothing has been sent. Review, then send from your SMS/email tool."
    summary["A2"].font = Font(italic=True, color="595959")
    summary.append([])

    high_risk = sum(1 for r in reminders if r["risk"] == "High")
    tiers = {t: sum(1 for m in rebooking if m["tier"] == t) for t in ("Due", "Overdue", "Lapsed")}
    value = sum(m["avg_spend"] for m in rebooking)
    lines = [("Appointment reminders", None),
             (f"  For appointments {('on ' + format(target, '%A, %b %d')) if target else '(none upcoming)'}", len(reminders)),
             ("  Firmer 'please confirm' version (repeat no-shows)", high_risk),
             ("  Not messaged (opted out / no contact info)", len(skipped_reminders)),
             ("Rebooking nudges", None),
             ("  Due (friendly nudge)", tiers["Due"]),
             ("  Overdue (usual gap exceeded)", tiers["Overdue"]),
             ("  Lapsed (win-back)", tiers["Lapsed"]),
             ("  Typical spend if they all rebook", float(value)),
             ("Clients not messaged, by reason", None)]
    for reason, count in sorted(skip_counts.items(), key=lambda kv: -kv[1]):
        lines.append((f"  {reason}", count))
    for label, value_ in lines:
        summary.append([label, value_])
        if value_ is None:
            summary.cell(summary.max_row, 1).font = Font(bold=True, size=12)
        elif isinstance(value_, float):
            summary.cell(summary.max_row, 2).number_format = '"$"#,##0.00'
    summary.column_dimensions["A"].width = 58
    summary.column_dimensions["B"].width = 14

    rows = [[r["start"].strftime("%a %b %d, %I:%M %p"), r["client_name"], r["barber"], r["service"],
             r["channel"], r["to"], r["risk"], r["message"], r["chars"], r["segments"]] for r in reminders]
    write_sheet(reminders_ws, [h for h, _ in REMINDER_COLUMNS], rows,
                [22, 22, 10, 18, 9, 28, 8, 80, 11, 9], wrap_cols=("Message",),
                empty_note="No appointments to remind.")
    for row in reminders_ws.iter_rows(min_row=2):
        if row[6].value == "High":
            for cell in row:
                cell.fill = RISK_FILL

    rows = [[m["tier"], m["client_id"], m["client_name"], m["channel"], m["to"], m["last_visit"],
             m["days_since"], m["usual_interval_days"], m["visits"], m["usual_barber"],
             m["usual_service"], m["avg_spend"], m["message"], m["chars"], m["segments"]] for m in rebooking]
    write_sheet(rebook_ws, [h for h, _ in REBOOK_COLUMNS], rows,
                [10, 10, 22, 9, 28, 12, 10, 12, 8, 12, 16, 11, 80, 11, 9],
                formats={"Last Visit": "yyyy-mm-dd", "Avg Spend": MONEY_FMT}, wrap_cols=("Message",),
                empty_note="No clients are due for a rebooking message.")
    for row in rebook_ws.iter_rows(min_row=2):
        fill = TIER_FILLS.get(row[0].value)
        if fill:
            for cell in row:
                cell.fill = fill

    rows = [[s["client_name"], s["start"].strftime("%a %b %d, %I:%M %p"), s["reason"]] for s in skipped_reminders]
    write_sheet(skipped_ws, ["Client", "Appointment", "Reason"], rows, [24, 22, 30],
                empty_note="Every reminder could be drafted.")
    for ws in (reminders_ws, rebook_ws):
        fit_landscape(ws)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# --------------------------------------------------------------------------
# Weekly report workbook
# --------------------------------------------------------------------------
def write_weekly_workbook(path: Path, ins: dict, issues: list, source_name: str, rows_used: int) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    daily_ws, services_ws, barbers_ws = wb.create_sheet("Daily"), wb.create_sheet("By Service"), wb.create_sheet("By Barber")
    busy_ws, noshow_ws, trend_ws, issues_ws = (wb.create_sheet("Busiest Hours"), wb.create_sheet("No-Shows"),
                                              wb.create_sheet("Weekly Trend"), wb.create_sheet("Data Issues"))
    start, end, cur, prev, has_prev = ins["start"], ins["end"], ins["current"], ins["previous"], ins["has_prev"]

    # ---- Summary ----
    summary["A1"] = f"Weekly Shop Report: {start:%a %b %d} to {end:%a %b %d, %Y}"
    summary["A1"].font = Font(bold=True, size=16)
    summary["A2"] = f"Source: {source_name}  |  Appointments used: {rows_used:,}  |  Rows with problems: {len(issues):,}"
    summary["A2"].font = Font(italic=True, color="595959")
    summary.append([])
    summary.append(["Metric", "This Week", "Last Week", "Change", "% Change"])
    for cell in summary[4]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    kpis = [("Revenue", "revenue", MONEY_FMT, True), ("Completed appointments", "completed", "#,##0", True),
            ("No-shows", "no_shows", "#,##0", False), ("No-show rate", "no_show_rate", "0.0%", False),
            ("Lost revenue (no-shows)", "lost_revenue", MONEY_FMT, False),
            ("Average ticket", "avg_ticket", MONEY_FMT, True), ("Cancellations", "cancelled", "#,##0", False),
            ("New clients", "new_clients", "#,##0", True), ("Returning clients", "returning_clients", "#,##0", True)]
    for label, key, fmt, up_is_good in kpis:
        row = summary.max_row + 1
        summary.cell(row, 1, label).font = Font(bold=True)
        summary.cell(row, 2, cur[key]).number_format = fmt
        if not has_prev:
            for col in (3, 4, 5):
                summary.cell(row, col, "n/a").alignment = Alignment(horizontal="right")
            continue
        summary.cell(row, 3, prev[key]).number_format = fmt
        diff = cur[key] - prev[key]
        summary.cell(row, 4, diff).number_format = "+0.0%;-0.0%;0.0%" if key == "no_show_rate" else fmt
        change = None if key == "no_show_rate" else pct_change(cur[key], prev[key])
        cell = summary.cell(row, 5, change if change is not None else "n/a")
        if change is None:
            cell.alignment = Alignment(horizontal="right")
        else:
            cell.number_format = CHANGE_FMT
        value = diff if key == "no_show_rate" else change
        if value:
            good = (value > 0) == up_is_good
            (summary.cell(row, 4) if key == "no_show_rate" else cell).font = GOOD_FONT if good else BAD_FONT

    row = summary.max_row + 2
    summary.cell(row, 1, "Highlights").font = Font(bold=True, size=13)
    for line in ins["highlights"]:
        row += 1
        summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = summary.cell(row, 1, "\u2022 " + line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        summary.row_dimensions[row].height = 16 * max(1, math.ceil(len(line) / 105))
    for idx, width in {1: 28, 2: 16, 3: 16, 4: 14, 5: 12}.items():
        summary.column_dimensions[get_column_letter(idx)].width = width

    # ---- detail sheets ----
    d = ins["daily"]
    write_sheet(daily_ws, list(d.columns), d.values.tolist(), [13, 7, 11, 11, 11, 13],
                formats={"Date": "yyyy-mm-dd", "Revenue": MONEY_FMT})

    def breakdown_sheet(ws, table, label):
        headers = [label, "Appointments", "Revenue", "% of Revenue", "Avg Price", "Last Week Revenue", "Change vs Last Week"]
        write_sheet(ws, headers, table.values.tolist(), [24, 14, 13, 13, 11, 18, 19],
                    formats={"Revenue": MONEY_FMT, "% of Revenue": "0.0%", "Avg Price": MONEY_FMT,
                             "Last Week Revenue": MONEY_FMT})
        for cells in ws.iter_rows(min_row=2, min_col=7, max_col=7):
            c = cells[0]
            if c.value is None:
                c.value, c.alignment = "n/a", Alignment(horizontal="right")
            else:
                c.number_format = CHANGE_FMT
                c.font = GOOD_FONT if c.value > 0 else BAD_FONT if c.value < 0 else Font()

    breakdown_sheet(services_ws, ins["services"], "Service")
    breakdown_sheet(barbers_ws, ins["barbers"], "Barber")

    grid = ins["grid"]
    busy_rows = [[hour] + grid.loc[hour].tolist() + [int(grid.loc[hour].sum())] for hour in grid.index]
    write_sheet(busy_ws, ["Start Hour"] + list(grid.columns) + ["All Days"], busy_rows,
                [14] + [9] * len(grid.columns) + [10])
    last_col = get_column_letter(len(grid.columns) + 1)
    busy_ws.conditional_formatting.add(
        f"B2:{last_col}{len(grid) + 1}",
        ColorScaleRule(start_type="min", start_color="FFFFFF", mid_type="percentile", mid_value=50,
                       mid_color="FFE699", end_type="max", end_color="F4B183"))
    busy_ws.cell(len(grid) + 3, 1, "Appointments (completed + no-show) by start time, last 8 weeks. Darker = busier.").font = Font(italic=True)

    ns = ins["no_shows"]
    write_sheet(noshow_ws, list(ns.columns), ns.values.tolist(), [24, 24, 10, 18, 14],
                formats={"Lost Revenue": MONEY_FMT}, empty_note="No no-shows this week.")
    t = ins["trend"]
    write_sheet(trend_ws, list(t.columns), t.values.tolist(), [15, 13, 14, 13],
                formats={"Revenue": MONEY_FMT, "No-Show Rate": "0.0%"})
    write_sheet(issues_ws, ["CSV Row", "Appointment", "Severity", "Issue"],
                [[i["row"], i["appointment_id"], i["severity"], i["issue"]] for i in issues],
                [9, 14, 11, 70], empty_note="No problems found in the input file.")

    # ---- charts ----
    if len(ins["services"]):
        bar = BarChart()
        bar.type = "bar"
        bar.title = "Revenue by service"
        bar.add_data(Reference(services_ws, min_col=3, min_row=1, max_row=len(ins["services"]) + 1), titles_from_data=True)
        bar.set_categories(Reference(services_ws, min_col=1, min_row=2, max_row=len(ins["services"]) + 1))
        bar.legend = None
        bar.y_axis.numFmt = '"$"#,##0'
        bar.x_axis.delete = False
        bar.y_axis.delete = False
        bar.x_axis.scaling.orientation = "maxMin"
        bar.height, bar.width = 7.5, 16
        summary.add_chart(bar, "G4")

    hours = BarChart()
    hours.type = "col"
    hours.title = "Appointments by start hour (last 8 weeks)"
    total_col = len(grid.columns) + 2
    hours.add_data(Reference(busy_ws, min_col=total_col, min_row=1, max_row=len(grid) + 1), titles_from_data=True)
    hours.set_categories(Reference(busy_ws, min_col=1, min_row=2, max_row=len(grid) + 1))
    hours.series[0].tx = SeriesLabel(v="Appointments")
    hours.legend = None
    hours.x_axis.delete = False
    hours.y_axis.delete = False
    hours.height, hours.width = 7.5, 16
    summary.add_chart(hours, "G20")

    fit_landscape(summary)
    summary.page_setup.fitToHeight = 1
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
